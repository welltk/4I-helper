import time
from tkinter import simpledialog, Toplevel
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import tkinter.font as tkFont
import webbrowser
from ttkbootstrap import Style
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Color, Font, Border, Side, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.workbook.properties import CalcProperties
from datetime import datetime
from copy import copy
import os
import json
import pandas as pd
import traceback
import math
import win32com.client
import threading
import pythoncom
from PIL import Image, ImageTk
from msoffcrypto.format.ooxml import OOXMLFile
from fouri_help import Help_Page, GifWindow, PngWindow

def rgb_to_str(rgb):
    if rgb is None:
        return None
    return f"{rgb}"

def str_to_rgb(s):
    if s is None or not s.startswith('FF'):
        return None
    return Color(rgb=s)

def load_styles_from_file(file_path):
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
            return data['styles'], data['column_dimensions'], data['row_dimensions']
    except FileNotFoundError:
        return {}, {}, {}
    except Exception as e:
        raise ValueError(f"스타일 로드 오류 발생 {file_path}: {e}")

def set_cell_style(cell, style):
    cell.font = Font(
        name=style['font']['name'],
        size=style['font']['size'],
        bold=style['font']['bold'],
        italic=style['font']['italic'],
        vertAlign=style['font']['vertAlign'],
        underline=style['font']['underline'],
        strike=style['font']['strike'],
        color=str_to_rgb(style['font']['color'])
    )
    cell.border = Border(
        left=Side(border_style=style['border']['left']['style'], color=str_to_rgb(style['border']['left']['color'])),
        right=Side(border_style=style['border']['right']['style'], color=str_to_rgb(style['border']['right']['color'])),
        top=Side(border_style=style['border']['top']['style'], color=str_to_rgb(style['border']['top']['color'])),
        bottom=Side(border_style=style['border']['bottom']['style'], color=str_to_rgb(style['border']['bottom']['color']))
    )
    cell.fill = PatternFill(
        fill_type=style['fill']['fill_type'],
        start_color=str_to_rgb(style['fill']['start_color']),
        end_color=str_to_rgb(style['fill']['end_color'])
    )
    cell.alignment = Alignment(
        horizontal=style['alignment']['horizontal'],
        vertical=style['alignment']['vertical'],
        text_rotation=style['alignment']['text_rotation'],
        wrap_text=style['alignment']['wrap_text'],
        shrink_to_fit=style['alignment']['shrink_to_fit'],
        indent=style['alignment']['indent']
    )
    cell.number_format = style['number_format']

def apply_styles(sheet, styles, column_dimensions, row_dimensions):
    for coord, style in styles.items():
        cell = sheet[coord]
        set_cell_style(cell, style)

    for col_letter, width in column_dimensions.items():
        sheet.column_dimensions[col_letter].width = width

    for row_idx, height in row_dimensions.items():
        if isinstance(row_idx, int):
            sheet.row_dimensions[row_idx].height = height

class GifWindow(Toplevel):
    def __init__(self, parent, gif_paths):
        super().__init__(parent)
        self.title("설명")
        self.geometry("1100x700")
        self.gif_paths = gif_paths
        self.current_gif_index = 0
        self.frames = []
        self.delays = []
        self.label = ttk.Label(self)
        self.label.pack()
        button_frame = ttk.Frame(self)
        button_frame.pack()
        self.prev_button = ttk.Button(button_frame, text="이전", command=self.prev_gif)
        self.prev_button.grid(row=0, column=0, padx=5, pady=5)
        self.next_button = ttk.Button(button_frame, text="다음", command=self.next_gif)
        self.next_button.grid(row=0, column=1, padx=5, pady=5)
        self.frame_index = 0
        self.current_timer = None
        self.load_first_frames()
        threading.Thread(target=self.load_remaining_frames).start()

    def load_first_frames(self):
        for gif_path in self.gif_paths:
            gif = Image.open(gif_path)
            frames = [ImageTk.PhotoImage(gif.copy())]
            delays = [gif.info['duration'] if 'duration' in gif.info else 100]
            if hasattr(gif, "is_animated") and gif.is_animated:
                try:
                    gif.seek(1)
                except EOFError:
                    pass
            self.frames.append(frames)
            self.delays.append(delays)
        self.update_frames()

    def load_remaining_frames(self):
        for i, gif_path in enumerate(self.gif_paths):
            gif = Image.open(gif_path)
            frames = self.frames[i]
            delays = self.delays[i]
            try:
                while True:
                    gif.seek(len(frames))
                    frames.append(ImageTk.PhotoImage(gif.copy()))
                    delays.append(gif.info['duration'] if 'duration' in gif.info else 100)
            except EOFError:
                pass

    def update_frames(self):
        if self.frames:
            frames = self.frames[self.current_gif_index]
            delays = self.delays[self.current_gif_index]
            self.label.config(image=frames[self.frame_index])
            delay = delays[self.frame_index]
            self.frame_index = (self.frame_index + 1) % len(frames)
            if self.current_timer:
                self.after_cancel(self.current_timer)
            self.current_timer = self.after(delay, self.update_frames)

    def prev_gif(self):
        self.current_gif_index = (self.current_gif_index - 1) % len(self.gif_paths)
        self.frame_index = 0
        self.update_frames()

    def next_gif(self):
        self.current_gif_index = (self.current_gif_index + 1) % len(self.gif_paths)
        self.frame_index = 0
        self.update_frames()

class LoadingDialog(Toplevel):
    def __init__(self, parent, title, message):
        super().__init__(parent)
        self.title(title)
        self.geometry("500x50")
        self.label = ttk.Label(self, text=message)
        self.label.pack(pady=10)
        self.protocol("WM_DELETE_WINDOW", self.on_closing)

    def on_closing(self):
        pass

class InfoDialog(tk.Toplevel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.title("정보")
        info_text = """오류 보고 및 피드백\n"""
        text_edit = tk.Text(self, wrap=tk.WORD)
        text_edit.insert(tk.END, info_text)
        text_edit.config(state=tk.DISABLED)
        text_edit.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

class NameJobDialog(tk.Toplevel):
    def __init__(self, parent=None, mappings=None):
        super().__init__(parent)
        self.parent = parent
        self.mappings = mappings if mappings is not None else {}
        self.title("직종 설정")
        self.geometry("500x500")
        style = ttk.Style()
        style.configure("Treeview", rowheight=30, font=('Noto Sans KR Medium', 14))
        style.configure("Treeview.Heading", font=('Noto Sans KR Medium', 16, 'bold'))
        self.table = ttk.Treeview(self, columns=("index", "name", "job"), show="headings")
        self.table.heading("index", text="번호")
        self.table.heading("name", text="이름")
        self.table.heading("job", text="직종")
        self.table.column("index", width=50, anchor='center')
        self.table.column("name", width=200, anchor='center')
        self.table.column("job", width=200, anchor='center')
        self.table.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        self.table.bind("<Double-1>", self.on_double_click)
        button_frame = ttk.Frame(self)
        button_frame.pack(fill=tk.X, padx=10, pady=10)
        self.add_button = ttk.Button(button_frame, text="추가", command=self.add_row)
        self.add_button.pack(side=tk.LEFT, padx=5)
        self.delete_button = ttk.Button(button_frame, text="삭제", command=self.delete_row)
        self.delete_button.pack(side=tk.LEFT, padx=5)
        self.apply_job_button = ttk.Button(button_frame, text="다중 직종 적용", command=self.apply_job_to_selected)
        self.apply_job_button.pack(side=tk.LEFT, padx=5)
        self.save_button = ttk.Button(button_frame, text="저장", command=self.save_and_close)
        self.save_button.pack(side=tk.RIGHT, padx=5)
        self.populate_table()

    def apply_job_to_selected(self):
        selected_items = self.table.selection()
        job = simpledialog.askstring("직종 선택", "적용할 직종을 입력하세요:")
        if job:
            for item in selected_items:
                self.table.set(item, "job", job)

    def populate_table(self):
        for idx, (name, job) in enumerate(self.mappings.items(), start=1):
            self.table.insert("", tk.END, values=(idx, name, job))

    def add_row(self):
        new_index = len(self.table.get_children()) + 1
        self.table.insert("", tk.END, values=(new_index, "", ""))

    def delete_row(self):
        selected_items = self.table.selection()
        for item in selected_items:
            self.table.delete(item)
        self.update_indices()

    def update_indices(self):
        for idx, item in enumerate(self.table.get_children(), start=1):
            values = self.table.item(item, "values")
            self.table.item(item, values=(idx, values[1], values[2]))

    def save_and_close(self):
        self.mappings.clear()
        for row in self.table.get_children():
            values = self.table.item(row, "values")
            self.mappings[values[1]] = values[2]
        self.parent.name_job_configuration = self.mappings
        self.parent.save_name_job_configuration_settings()
        messagebox.showinfo("직종 설정", "직종 설정 저장 완료.")
        self.destroy()

    def on_double_click(self, event):
        selected_item = self.table.selection()
        if not selected_item:
            return
        item = selected_item[0]
        column = self.table.identify_column(event.x)
        row = self.table.identify_row(event.y)
        col_idx = int(column.replace('#', '')) - 1
        current_value = self.table.item(item, "values")[col_idx]
        x, y, width, height = self.table.bbox(item, column)
        height = 40
        entry = tk.Text(self.table, height=height // 20, width=width // 8, font=("Noto Sans KR Medium", 14))
        entry.place(x=x, y=y, anchor='nw', width=width, height=height)
        entry.insert(tk.END, current_value)

        def on_entry_confirm(event=None):
            new_value = entry.get("1.0", tk.END).strip()
            current_values = list(self.table.item(item, "values"))
            current_values[col_idx] = new_value
            self.table.item(item, values=current_values)
            entry.destroy()

        entry.bind("<Return>", on_entry_confirm)
        entry.bind("<FocusOut>", on_entry_confirm)

    def edit_cell(self, row, col_idx):
        bbox = self.table.bbox(row, f'#{col_idx+1}')
        if not bbox:
            return
        x, y, width, height = bbox
        entry = tk.Entry(self.table)
        entry.place(x=x, y=y, width=width, height=height)
        value = self.table.item(row, 'values')[col_idx]
        entry.insert(0, value)
        entry.focus()

        def save_value(event=None):
            new_value = entry.get()
            values = list(self.table.item(row, 'values'))
            values[col_idx] = new_value
            self.table.item(row, values=values)
            entry.destroy()

        entry.bind('<Return>', save_value)
        entry.bind('<FocusOut>', save_value)

class PasswordDialog(tk.Toplevel):
    def __init__(self, parent, callback):
        super().__init__(parent)
        self.title("비밀번호 입력")
        self.geometry("600x150")
        self.callback = callback
        self.label = ttk.Label(self, text="엑셀 파일에 설정할 비밀번호를 입력하세요 (영어와 숫자만 입력):")
        self.label.pack(pady=10)
        self.password_entry = ttk.Entry(self, show='*')
        self.password_entry.pack(pady=5)
        self.password_entry.focus_set()
        self.password_entry.bind('<Return>', self.submit)
        self.submit_button = ttk.Button(self, text="확인", command=self.submit)
        self.submit_button.pack(pady=5)
        self.protocol("WM_DELETE_WINDOW", self.on_closing)

    def is_valid_password(self, password):
        return all(char.isalnum() for char in password)

    def submit(self, event=None):
        password = self.password_entry.get()
        if not self.is_valid_password(password):
            messagebox.showwarning("유효성 검사 실패", "비밀번호는 영어와 숫자만 포함해야 합니다.")
            return
        self.callback(password)
        self.destroy()

    def on_closing(self):
        self.callback(None)
        self.destroy()

class ErrorDialog(tk.Toplevel):
    def __init__(self, parent=None, error_message=""):
        super().__init__(parent)
        self.title("오류")
        self.geometry("500x500")
        layout = ttk.Frame(self)
        layout.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        label = ttk.Label(layout, text="오류가 발생했습니다. 아래 내용을 복사하여 제작자에게 문의하세요:")
        label.pack(pady=10)
        self.text_edit = tk.Text(layout, wrap=tk.WORD)
        self.text_edit.insert(tk.END, error_message)
        self.text_edit.config(state=tk.DISABLED)
        self.text_edit.pack(fill=tk.BOTH, expand=True, pady=10)
        close_button = ttk.Button(layout, text="닫기", command=self.destroy)
        close_button.pack()

class merge(ttk.Frame):
    def __init__(self, parent=None, mappings=None):
        super().__init__(parent)
        if mappings is None:
            mappings = {}
        self.settings = None
        self.name_job_configuration = None
        self.mappings = mappings
        self.original_files = []
        self.PS_health_insurance_file = ''
        self.health_insurance_file = ''
        self.np_insurance_file = ''
        self.ep_insurance_file = ''
        self.ia_insurance_file = ''
        self.include_income_data = tk.BooleanVar()
        self.add_part_time = tk.BooleanVar()
        self.init_ui()

    def apply_conditional_formatting(self, sheet):
        colors = {
            "기타급여1": "9ba1db",
            "기타급여4": "92d6ca",
            "그외 직종": "a6a7d8",
            "행정대체": "ffceb0",
            "기간제교사": "ffef99"
        }
        for category, color in colors.items():
            fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            formula = f'$C7="{category}"'
            rule = FormulaRule(formula=[formula], fill=fill)
            sheet.conditional_formatting.add('C7:C100', rule)

    def apply_conditional_formatting_zero(self, sheet, cell_ranges):
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        purple_fill = PatternFill(start_color="8A88D1", end_color="8A88D1", fill_type="solid")
        yellow_font = Font(color="FFFF00")
        for cell_range in cell_ranges:
            sheet.conditional_formatting.add(cell_range, CellIsRule(operator='greaterThan', formula=['0'], fill=red_fill, font=yellow_font))
            sheet.conditional_formatting.add(cell_range, CellIsRule(operator='lessThan', formula=['0'], fill=purple_fill, font=yellow_font))

    def open_youtube(self, url):
        webbrowser.open(url)

    def change_cursor(self, event):
        widget = event.widget
        tag_indices = widget.tag_names(tk.CURRENT)
        if any(tag in self.link_positions for tag in tag_indices):
            widget.config(cursor="hand2")
        else:
            widget.config(cursor="arrow")

    def init_ui(self):
        self.font_path = "font/Noto Sans KR Medium.ttf"
        style = Style(theme='flatly')
        style.configure("Custom.TButton", font=("Noto Sans KR Medium", 12), background="#f6fafa", foreground="#000000", bordercolor="#365486", borderwidth=0.5, focusthickness=0.5, focuscolor='none')
        self.button_font = tkFont.Font(family="Noto Sans KR Medium", size=13)
        tkFont.nametofont("TkDefaultFont").configure(family="Noto Sans KR Medium", size=13)
        self.style = Style(theme='flatly')
        self.style.configure('Custom.TButton', font=(self.font_path, 13))
        self.style.configure('.', background='white', foreground='black')
        self.style.configure("TFrame", background="white")
        central_frame = ttk.Frame(self)
        central_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        layout = ttk.Frame(central_frame)
        layout.pack(fill=tk.BOTH, expand=True)
        layout.pack_propagate(0)
        layout.columnconfigure(0, weight=1)
        self.source_file_paths = ttk.Entry(layout, state='readonly')
        self.target_file_path = ttk.Entry(layout, state='readonly')
        source_file_button = self.create_button(layout, '불러오기', self.select_files)
        target_file_button = self.create_button(layout, '불러오기', self.select_target_file)
        self.text = tk.Text(layout, wrap="word", font=("Noto Sans KR Medium", 15), height=1, width=40, bd=0, highlightthickness=0, padx=0, pady=0)
        self.text.grid(row=2, column=0, pady=(40, 6), sticky='w')
        self.text.insert("1.0", "나이스 급여대장 및 보험 고지내역 파일")
        self.text.tag_add("link1", "1.0", "1.8")
        self.text.tag_add("link2", "1.11", "1.18")
        self.text.tag_config("link1", foreground="blue", underline=True)
        self.text.tag_config("link2", foreground="blue", underline=True)
        self.text.config(state="disabled", cursor="arrow")
        self.text_2 = tk.Text(layout, wrap="word", font=("Noto Sans KR Medium", 15), height=1, width=40, bd=0, highlightthickness=0, padx=0, pady=0)
        self.text_2.grid(row=4, column=0, pady=(10, 6), sticky='w')
        self.text_2.insert("1.0", "엑셀 파일")
        self.text_2.tag_add("link3", "1.0", "1.29")
        self.text_2.tag_config("link3", foreground="blue", underline=True)
        self.text_2.config(state="disabled", cursor="arrow")
        self.link_positions = {
            "link1": (["doc/gif/salary1.gif", "doc/gif/salary2.gif", "doc/gif/salary3.gif"]),
            "link2": None,
        }
        self.open_youtube = {
            "link3": ("https://blog.naver.com/brojun88/")
        }

        class OptionDialog(tk.Toplevel):
            def __init__(self, parent, title, options, callback):
                super().__init__(parent)
                self.title(title)
                self.geometry("300x200")
                self.callback = callback
                self.var = tk.StringVar(value=options[0])
                for option in options:
                    rb = ttk.Radiobutton(self, text=option, variable=self.var, value=option)
                    rb.pack(anchor=tk.W, padx=10, pady=5)
                confirm_button = ttk.Button(self, text="확인", command=self.confirm)
                confirm_button.pack(pady=10)

            def confirm(self):
                self.callback(self.var.get())
                self.destroy()

        def text_click(event):
            widget = event.widget
            tag_indices = widget.tag_names(tk.CURRENT)
            for tag in tag_indices:
                if tag in self.link_positions:
                    if tag == 'link2':
                        options = ["건강보험", "국민연금", "고용산재"]
                        def handle_selection(selected_option):
                            if selected_option == "건강보험":
                                png_paths = ["doc/png/health1.png", "doc/png/health2.png", "doc/png/health3.png", "doc/png/health4.png"]
                            elif selected_option == "국민연금":
                                png_paths = ["doc/png/pension1.png", "doc/png/pension2.png", "doc/png/pension3.png"]
                            elif selected_option == "고용산재":
                                png_paths = ["doc/png/emp1.png", "doc/png/emp2.png", "doc/png/emp3.png", "doc/png/emp4.png"]
                            PngWindow(self, png_paths)
                        OptionDialog(self, "보험 항목 선택", options, handle_selection)
                    else:
                        gif_paths = self.link_positions[tag]
                        GifWindow(self, gif_paths)
                    break
                elif tag in self.open_youtube:
                    url = self.open_youtube[tag]
                    webbrowser.open_new(url)
                    break

        self.text.bind("<Button-1>", text_click)
        self.text.bind("<Motion>", self.change_cursor)
        self.text_2.bind("<Button-1>", text_click)
        self.text_2.bind("<Motion>", self.change_cursor)
        self.source_file_paths.grid(row=3, column=0, padx=5, pady=(10, 10), sticky='we')
        source_file_button.grid(row=3, column=1, padx=5, pady=(10, 10))
        self.target_file_path.grid(row=5, column=0, padx=5, pady=(10, 10), sticky='we')
        target_file_button.grid(row=5, column=1, padx=5, pady=(10, 10))
        self.target_file_path.grid(row=5, column=0, padx=5, pady=(10, 10), sticky='we')
        target_file_button.grid(row=5, column=1, padx=5, pady=(10, 10))
        self.include_income_checkbutton = ttk.Checkbutton(layout, text="일용근로자 (대체직, 시간강사 등) 고용보험 개인부담금 급여대장에 자동 추가", variable=self.include_income_data)
        self.include_income_checkbutton.grid(row=6, column=0, sticky='w')
        button_frame = ttk.Frame(layout)
        button_frame.grid(row=9, column=0, columnspan=2, pady=(10, 20))
        for i in range(5):
            button_frame.columnconfigure(i, weight=1)
        self.create_button(button_frame, '직종 설정', self.open_name_job_dialog).grid(row=0, column=0, padx=5, pady=5)
        self.create_button(button_frame, '설정 저장하기', self.save_settings).grid(row=0, column=1, padx=5, pady=5)
        self.create_button(button_frame, '설정 불러오기', self.load_settings).grid(row=0, column=2, padx=5, pady=5)
        self.create_button_process(button_frame, '자료 병합하기', self.process_files_in_thread).grid(row=0, column=3, padx=5, pady=5)
        self.create_button(button_frame, '정보', self.open_info_dialog).grid(row=0, column=4, padx=5, pady=5)
        blog_label = ttk.Label(layout, text="파란 글자를 누르면 설명을 볼 수 있습니다.", font=("Noto Sans KR Medium", 15))
        blog_label.grid(row=10, column=0, columnspan=2, pady=(10, 5))

    def create_button(self, parent, text, command):
        style = ttk.Style()
        style.configure('Custom.TButton', font=("Noto Sans KR Medium", 12))
        button = ttk.Button(parent, text=text, command=command, style='Custom.TButton')
        return button

    def create_button_process(self, parent, text, command):
        style = ttk.Style()
        style.configure('Custom.TButton', font=("Noto Sans KR Medium", 12))
        button = ttk.Button(parent, text=text, command=command, style="info")
        return button

    def create_hyperlink(self, parent, text, url, font=None):
        def callback(event):
            webbrowser.open_new(url)
        label = ttk.Label(parent, text=text, foreground="blue", cursor="hand2", font=font)
        label.bind("<Button-1>", callback)
        return label

    def select_files(self):
        files = filedialog.askopenfilenames(title="파일 선택", filetypes=(("All Files", "*.*"),))
        if files:
            self.source_file_paths.config(state='normal')
            self.source_file_paths.delete(0, tk.END)
            self.source_file_paths.insert(0, '; '.join(files))
            self.source_file_paths.config(state='readonly')
            self.handle_file_selection()

    def select_target_file(self):
        file = filedialog.askopenfilename(title="계산기 파일 선택", filetypes=(("All Files", "*.*"),))
        if file:
            self.target_file_path.config(state='normal')
            self.target_file_path.delete(0, tk.END)
            self.target_file_path.insert(0, file)
            self.target_file_path.config(state='readonly')

    def handle_file_selection(self):
        files = self.source_file_paths.get().split('; ')
        self.categorize_files(files)

    def categorize_files(self, files):
        self.original_files = []
        self.PS_health_insurance_file = ''
        self.health_insurance_file = ''
        self.np_insurance_file = ''
        self.ep_insurance_file = ''
        self.ia_insurance_file = ''
        for file in files:
            try:
                df = pd.read_excel(file, nrows=2)
                if '기관분류' in df.columns and '소득구분' in df.columns:
                    self.original_files.append(file)
                elif '사업장관리번호' in df.columns:
                    if '당월분_기준소득월액' not in df.columns:
                        if str(df.iloc[0]['사업장관리번호'])[-1] == '2':
                            self.PS_health_insurance_file = file
                        else:
                            self.health_insurance_file = file
                    else:
                        self.np_insurance_file = file
                elif '근로자실업급여보험료' in df.columns or '근로자실업급여보험료' in df.iloc[0].tolist():
                    self.ep_insurance_file = file
                elif '보험료합계(①+②+③)' in df.columns and '근로자실업급여보험료' not in df.columns:
                    self.ia_insurance_file = file
            except Exception as e:
                print(f"파일 분류 오류 발생 {file}: {e}")

    def open_name_job_dialog(self):
        dialog = NameJobDialog(self, self.name_job_configuration)
        self.wait_window(dialog)

    def open_info_dialog(self):
        dialog = InfoDialog(self)
        self.wait_window(dialog)

    def save_settings(self):
        try:
            settings = {
                '원본파일들': self.source_file_paths.get().split('; '),
                '계산기_파일': self.target_file_path.get(),
                'include_income_data': self.include_income_data.get(),
                'add_part_time': self.add_part_time.get()
            }
            with open('Settings_Merger.json', 'w', encoding='utf-8') as f:
                json.dump(settings, f, ensure_ascii=False, indent=4)
            messagebox.showinfo('설정', '설정 저장 완료.')
        except Exception as e:
            print(f"설정 저장에 오류가 발생함: {e}")

    def save_name_job_configuration_settings(self):
        try:
            with open('name_job_configuration.json', 'w') as f:
                json.dump(self.name_job_configuration, f)
        except Exception as e:
            print(f"직종 설정 저장에 오류가 발생함: {e}")

    def load_settings(self):
        try:
            with open('Settings_Merger.json', 'r', encoding='utf-8') as f:
                settings = json.load(f)
            self.source_file_paths.config(state='normal')
            self.source_file_paths.delete(0, tk.END)
            self.source_file_paths.insert(0, '; '.join(settings.get('원본파일들', [])))
            self.source_file_paths.config(state='readonly')
            self.target_file_path.config(state='normal')
            self.target_file_path.delete(0, tk.END)
            self.target_file_path.insert(0, settings.get('계산기_파일', ''))
            self.target_file_path.config(state='readonly')
            self.include_income_data.set(settings.get('include_income_data', False))
            self.add_part_time.set(settings.get('add_part_time', False))
            self.load_name_job_configuration_settings()
            self.handle_file_selection()
            messagebox.showinfo('설정', '설정 불러오기 완료.')
        except FileNotFoundError:
            pass
        except json.JSONDecodeError as e:
            messagebox.showwarning(self, '설정', f'JSON 디코드 오류: {e}')
        except Exception as e:
            messagebox.showwarning(self, '설정', f'오류 발생: {e}')

    def load_name_job_configuration_settings(self):
        try:
            if os.path.exists('name_job_configuration.json'):
                with open('name_job_configuration.json', 'r') as f:
                    self.name_job_configuration = json.load(f)
        except Exception as e:
            print(f"직종 설정 불러오기에 오류가 발생함: {e}")

    def find_next_empty_row(self, sheet, start_col=1, start_row=2):
        while sheet.cell(row=start_row, column=start_col).value:
            start_row += 1
        return start_row

    def copy_format_row(self, sheet, source_row_idx, target_row_idx):
        for col in range(1, sheet.max_column + 1):
            source_cell = sheet.cell(row=source_row_idx, column=col)
            target_cell = sheet.cell(row=target_row_idx, column=col)
            target_cell.value = source_cell.value
            if source_cell.has_style:
                target_cell.font = copy(source_cell.font)
                target_cell.border = copy(source_cell.border)
                target_cell.fill = copy(source_cell.fill)
                target_cell.number_format = copy(source_cell.number_format)
                target_cell.protection = copy(source_cell.protection)
                target_cell.alignment = copy(source_cell.alignment)

    def get_next_month(self, date):
        year = date.year + (date.month // 12)
        month = date.month % 12 + 1
        return datetime(year, month, 1)

    def merge_cells(self, sheet, start_row, end_row, start_col, end_col):
        sheet.merge_cells(start_row=start_row, start_column=start_col, end_row=end_row, end_column=end_col)

    def adjust_column_width(self, sheet, column_widths):
        for col_letter, width in column_widths.items():
            sheet.column_dimensions[col_letter].width = width

    def process_files_in_thread(self):
        def process():
            try:
                pythoncom.CoInitialize()
                self.process_files()
                pythoncom.CoUninitialize()
            except Exception as e:
                error_message = f'오류가 발생하였습니다.: {e}\n{traceback.format_exc()}'
                self.show_error_message('오류', error_message)
                pythoncom.CoUninitialize()
            finally:
                self.loading_dialog.destroy()

        self.loading_dialog = LoadingDialog(self, "병합 중...", "파일을 병합하는 중입니다. 잠시만 기다려주세요...")
        threading.Thread(target=process).start()

    def show_info_message(self, title, message):
        messagebox.showinfo(title, message)

    def show_error_message(self, title, message):
        messagebox.showerror(title, message)

    def set_password_on_excel(self, file_path, password):
        temp_file = 'temp_' + os.path.basename(file_path)
        with open(file_path, 'rb') as plain_file:
            ooxml_file = OOXMLFile(plain_file)
            with open(temp_file, 'wb') as encrypted_file:
                ooxml_file.encrypt(password, encrypted_file)
        os.replace(temp_file, file_path)

    def run_excel_macro(self, file_path):
        try:
            pythoncom.CoInitialize()
            excel = win32com.client.Dispatch("HCell.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            xlCalculationManual = -4135
            workbook = excel.Workbooks.Open(file_path)
            excel.Calculation = xlCalculationManual
            excel.CalculateFullRebuild()
            xlCalculationAutomatic = -4105
            excel.Calculation = xlCalculationAutomatic
            new_tgt_file = file_path.replace('.xlsm', f'.xlsm')
            workbook.SaveAs(new_tgt_file)
            workbook.Close(False)
            del workbook
            excel.Quit()
            del excel
            pythoncom.CoUninitialize()
            return new_tgt_file
        except Exception as e:
            if 'excel' in locals():
                excel.Quit()
            pythoncom.CoUninitialize()
            raise

    def process_files(self):
        src_files = self.source_file_paths.get().split('; ')
        PS_health_insurance_file = self.PS_health_insurance_file
        health_insurance_file = self.health_insurance_file
        np_insurance_file = self.np_insurance_file
        ep_insurance_file = self.ep_insurance_file
        ia_insurance_file = self.ia_insurance_file
        tgt_file = self.target_file_path.get()
        if not tgt_file:
            messagebox.showwarning('경고', '계산기 파일을 선택하세요.')
            return
        try:
            if not os.path.exists(tgt_file):
                raise FileNotFoundError(f"파일을 찾을 수 없습니다: {tgt_file}")
            wb = load_workbook(tgt_file, read_only=True, keep_vba=True, data_only=False)
            temp_wb = Workbook()
            temp_wb.remove(temp_wb.active)
            for sheet_name in wb.sheetnames:
                temp_sheet = temp_wb.create_sheet(sheet_name)
                original_sheet = wb[sheet_name]
                for row in original_sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            temp_sheet[cell.coordinate].value = cell.value
            temp_wb.save("temp_file.xlsm")
            wb.close()
            wb = load_workbook("temp_file.xlsm")
            new_wb = Workbook()
            new_wb.remove(new_wb.active)
            new_wb.properties.calcPr = CalcProperties(calcId=124519, fullCalcOnLoad=True)
            for sheet_name in wb.sheetnames:
                original_sheet = wb[sheet_name]
                new_sheet = new_wb.create_sheet(sheet_name)
                for row in original_sheet.iter_rows():
                    for cell in row:
                        new_sheet[cell.coordinate].value = cell.value
            styles, column_dimensions, row_dimensions = load_styles_from_file('styles.json')
            for sheet_name, sheet_styles in styles.items():
                if sheet_name in new_wb.sheetnames:
                    apply_styles(new_wb[sheet_name], sheet_styles, column_dimensions[sheet_name], row_dimensions[sheet_name])
            self_seperate_transfer_sheet = new_wb['2. 본인부담금 별도이체']
            self_seperate_transfer_sheet.merge_cells('A18:H18')
            self_seperate_transfer_sheet.merge_cells('A25:E25')
            total_i_sheet = new_wb['1. 4대보험 부담금 합계']
            for row in range(82, 96):
                total_i_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            for row in range(1, 5):
                self_seperate_transfer_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            institution_request_sheet = new_wb['3. 기관부담금 신청(자료집계)']
            institution_request_sheet.merge_cells('A1:I1')
            institution_seperate_transfer_sheet = new_wb['4. 기관부담금 세외이체 합계']
            institution_seperate_transfer_sheet.merge_cells('A1:I1')
            final_details_sheet = new_wb['5. 최종 납부내역']
            final_details_sheet.merge_cells('A38:C38')
            final_pay_sheet = new_wb['7. 최종 반환용 합계']
            final_pay_sheet.merge_cells('A2:C2')
            for row in range(39, 46):
                final_details_sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
            for row in range(29, 35):
                final_details_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            column_widths = {'A': 15}
            for col in range(2, 11):
                final_details_sheet.column_dimensions[get_column_letter(col)].width = 20
            for col, width in column_widths.items():
                final_details_sheet.column_dimensions[col].width = width
            ps_health_sheet = new_wb['0. 공무원 건강보험 부담금']
            current_date = datetime.now()
            next_month_date = self.get_next_month(current_date)
            ps_health_sheet['M3'] = f"{current_date.year}년 {current_date.month}월"
            ps_health_sheet['M4'] = f"{next_month_date.year}년 {next_month_date.month}월"
            formatted_date = f'{current_date.year}-{next_month_date.month:02}-08'
            Current_formatted_date = f'{current_date.year}-{current_date.month:02}-17'
            Current2_formatted_date = f'{current_date.year}-{current_date.month:02}-29'
            rows_to_update = [9, 11, 12, 13, 14]
            for row in rows_to_update:
                cell = f'C{row}'
                final_details_sheet[cell] = formatted_date
            Current_rows_to_update = [20, 26]
            for row in Current_rows_to_update:
                cell = f'C{row}'
                final_details_sheet[cell] = Current_formatted_date
            final_details_sheet['C21'] = Current2_formatted_date
            columns = [
                '성명', '직종', '건강보험', '건강보험연말정산', '건강보험퇴직정산',
                '건강보험복직정산', '노인장기요양보험', '장기요양연말정산', '장기요양퇴직정산',
                '장기요양복직정산', '국민연금', '고용보험', '고용보험연말정산', '고용보험퇴직정산'
            ]
            HI_columns = [
                '성명', '고지금액', '요양고지보험료', '건강환급금이자', '요양환급금이자', '가입자총납부할보험료'
            ]
            NP_columns = [
                '성명', '총부담금계_(사용자부담금)', '총부담금계_(본인기여금)'
            ]
            EP_columns = [
                '근로자명', '근로자실업급여보험료', '사업주실업급여보험료', '사업주고안직능보험료', '합계'
            ]
            IA_columns = [
                '근로자명', '산정보험료(해당월①)', '재산정보험료(해당년도②)', '정산보수총액', '정산보험료(③)', '보험료합계(①+②+③)'
            ]
            sheets = {}
            for sheet_name in wb.sheetnames:
                if "급여" in sheet_name:
                    sheets['salary'] = wb[sheet_name]
                elif "공무원 건강보험" in sheet_name:
                    sheets['PS_Health'] = wb[sheet_name]
                    sheets['PS_HI_edi'] = wb[sheet_name]
                elif "4대보험" in sheet_name:
                    sheets['Total_I'] = wb[sheet_name]
                elif "비공무원 건강보험" in sheet_name:
                    sheets['HI_edi'] = wb[sheet_name]
                elif "국민연금" in sheet_name:
                    sheets['NP_edi'] = wb[sheet_name]
                elif "고용보험" in sheet_name:
                    sheets['EP_edi'] = wb[sheet_name]
                elif "산재보험" in sheet_name:
                    sheets['IA_edi'] = wb[sheet_name]
                elif "별도이체" in sheet_name:
                    sheets['Etc_fee'] = wb[sheet_name]
                elif "최종" in sheet_name:
                    sheets['Final_details'] = wb[sheet_name]
            self.tgt_row_idx_salary = self.find_next_empty_row(self.sheets['salary'])
            self.tgt_row_idx_Total_I = self.find_next_empty_row(self.sheets['Total_I'], start_row=7)
            self.tgt_row_idx_PS_HI_edi = self.find_next_empty_row(self.sheets['PS_HI_edi'])
            self.tgt_row_idx_HI_edi = self.find_next_empty_row(self.sheets['HI_edi'])
            self.tgt_row_idx_NP_edi = self.find_next_empty_row(self.sheets['NP_edi'])
            self.tgt_row_idx_EP_edi = self.find_next_empty_row(self.sheets['EP_edi'])
            self.tgt_row_idx_IA_edi = self.find_next_empty_row(self.sheets['IA_edi'])

            def convert_to_number(value):
                if isinstance(value, (int, float)):
                    return value
                if isinstance(value, str):
                    value = value.replace(',', '')
                    try:
                        return float(value)
                    except ValueError:
                        return value
                return value

            salary_names = set()
            for row in self.sheets['salary'].iter_rows(min_row=2, max_col=1, max_row=self.tgt_row_idx_salary - 1):
                for cell in row:
                    if cell.value:
                        salary_names.add(cell.value)

            def extract_birthdate(value, is_rrn=False):
                if is_rrn:
                    if isinstance(value, str) and len(value) >= 6:
                        return value[:6]
                else:
                    if isinstance(value, str):
                        return value.replace('-', '')
                return None

            def copy_data(df, is_original_file=False):
                tgt_row_idx_salary = self.tgt_row_idx_salary
                tgt_row_idx_Total_I = self.tgt_row_idx_Total_I
                row_idx_PS_Health = self.find_next_empty_row(self.sheets['PS_Health'], start_row=7)
                existing_columns = set(df.columns)
                tgt_columns = {cell.value: cell.column for cell in self.sheets['salary'][1]}
                salary_columns = [cell.value for cell in self.sheets['salary'][1]]
                processed_individuals = set()
                columns_to_check = [
                    '건강보험', '건강보험연말정산', '건강보험퇴직정산', '건강보험복직정산',
                    '노인장기요양보험', '장기요양연말정산', '장기요양퇴직정산', '장기요양복직정산',
                    '국민연금', '고용보험', '고용보험연말정산', '고용보험퇴직정산'
                ]
                if is_original_file:
                    if '성명' not in df.columns or '개인번호' not in df.columns:
                        raise KeyError("원본 파일에 '성명' 또는 '개인번호' 열이 없습니다.")
                    for idx, row in df.dropna(subset=['성명', '개인번호']).iterrows():
                        identifier = (row['성명'], row['개인번호'])
                        if identifier in processed_individuals:
                            continue
                        has_insurance_values = any(not pd.isna(row.get(col_name, None)) and row.get(col_name, None) != "" for col_name in columns_to_check)
                        if not has_insurance_values:
                            continue
                name_count_map = {}
                for idx, row in df.dropna(subset=['성명']).iterrows():
                    name = row['성명']
                    birthdate = None
                    if '주민등록번호' in row:
                        birthdate = extract_birthdate(row['주민등록번호'], is_rrn=True)
                    elif '생년월일' in row:
                        birthdate = extract_birthdate(row['생년월일'])
                    if name in name_count_map:
                        name_count_map[name] += 1
                        unique_name = f"{name}{name_count_map[name]}"
                    else:
                        name_count_map[name] = 1
                        unique_name = name
                    if tgt_row_idx_salary > 21:
                        self.copy_format_row(self.sheets['salary'], 21, tgt_row_idx_salary)
                    for col_name in salary_columns:
                        col_idx = tgt_columns[col_name]
                        if col_name in existing_columns:
                            value = row.get(col_name, 0)
                            if pd.isna(value) or value == "" or value == "None":
                                value = 0
                            if col_name == '직종':
                                if '시도직종' in df.columns and not any(keyword in row['직종'] for keyword in ['기간제', '계약제', '계약직']):
                                    value = row.get('시도직종', '')
                                else:
                                    value = row.get('직종', '')
                                if not value:
                                    value = self.name_job_configuration.get(name, '기타')
                            self.sheets['salary'].cell(row=tgt_row_idx_salary, column=col_idx, value=value)
                            if value is not None and isinstance(value, (int, float)) and col_name != '직종':
                                self.sheets['salary'].cell(row=tgt_row_idx_salary, column=col_idx).number_format = '#,##0_);-#,##0'
                        else:
                            if col_name not in ['성명', '직종']:
                                self.sheets['salary'].cell(row=tgt_row_idx_salary, column=col_idx, value=0)
                                self.sheets['salary'].cell(row=tgt_row_idx_salary, column=col_idx).number_format = '#,##0_);-#,##0'
                    self.sheets['salary'].cell(row=tgt_row_idx_salary, column=tgt_columns['성명'], value=unique_name)
                    tgt_row_idx_salary += 1
                    if (
                        any(keyword in row['직종'] for keyword in ['교원', '일반', '지방', '행정']) and
                        not ('지방' in row['직종'] and '대체' in row['직종']) and
                        not ('교원' in row['직종'] and any(exclude_keyword in row['직종'] for exclude_keyword in ['기간', '계약'])) and
                        not ('행정' in row['직종'] and '대체' in row['직종'])
                    ):
                        self.sheets['PS_Health'].cell(row=row_idx_PS_Health, column=1, value=unique_name)
                        row_idx_PS_Health += 1
                    elif any(keyword in row['직종'] for keyword in ['계약', '기간', '대체', '공무직']):
                        job_value = '기간제교사' if any(keyword in row['직종'] for keyword in ['계약', '기간']) else row.get('시도직종', row['직종'])
                        self.sheets['Total_I'].cell(row=tgt_row_idx_Total_I, column=1, value=unique_name)
                        self.sheets['Total_I'].cell(row=tgt_row_idx_Total_I, column=2, value=job_value)
                        tgt_row_idx_Total_I += 1
                    self.tgt_row_idx_salary = tgt_row_idx_salary
                    self.tgt_row_idx_Total_I = tgt_row_idx_Total_I

            def copy_income_data(df):
                tgt_row_idx_salary = self.tgt_row_idx_salary
                tgt_row_idx_Total_I = self.tgt_row_idx_Total_I
                tgt_columns = {cell.value: cell.column for cell in self.sheets['salary'][1]}
                salary_columns = {cell.value for cell in self.sheets['salary'][1]}
                total_i_columns = {cell.value: cell.column for cell in self.sheets['Total_I'][6]}
                for _, row in df.iterrows():
                    if row[1] == '일용근로소득':
                        if tgt_row_idx_salary > 21:
                            self.copy_format_row(self.sheets['salary'], 21, tgt_row_idx_salary)
                        self.sheets['salary'].cell(row=tgt_row_idx_salary, column=tgt_columns['성명'], value=row[3])
                        try:
                            value = float(row[7]) if row[7] != "" else None
                            if value is not None:
                                value = math.floor(value * 0.0009)
                                cell = self.sheets['salary'].cell(row=tgt_row_idx_salary, column=tgt_columns['고용보험'], value=value * 10)
                                cell.number_format = '#,##0_);[Magenta]("₩"#,##0)'
                            else:
                                self.sheets['salary'].cell(row=tgt_row_idx_salary, column=tgt_columns['고용보험'], value=None)
                        except ValueError:
                            self.sheets['salary'].cell(row=tgt_row_idx_salary, column=tgt_columns['고용보험'], value=None)
                        for col_name in ['건강보험', '국민연금', '노인장기요양보험', '고용보험연말정산']:
                            if col_name in salary_columns:
                                self.sheets['salary'].cell(row=tgt_row_idx_salary, column=tgt_columns[col_name], value=0)
                        if row[3] in self.name_job_configuration:
                            job_value = self.name_job_configuration[row[3]]
                            if job_value == '시간강사':
                                job_value = '시간강사(대체)'
                        else:
                            job_value = '기타'
                        self.sheets['salary'].cell(row=tgt_row_idx_salary, column=tgt_columns['직종'], value=job_value)
                        if '성명' in total_i_columns:
                            self.sheets['Total_I'].cell(row=tgt_row_idx_Total_I, column=total_i_columns['성명'], value=row[3])
                        tgt_row_idx_salary += 1
                        tgt_row_idx_Total_I += 1
                self.tgt_row_idx_salary = tgt_row_idx_salary
                self.tgt_row_idx_Total_I = tgt_row_idx_Total_I

            def copy_PS_health_insurance_data(df):
                tgt_row_idx_PS_HI_edi = self.tgt_row_idx_PS_HI_edi
                tgt_columns = {cell.value: cell.column for cell in self.sheets['PS_HI_edi'][1]}
                processed_individuals = set()
                name_count_map = {}
                for _, row in df.iterrows():
                    if not pd.isna(row['성명']) and '주민등록번호' in row:
                        birthdate = row['주민등록번호'][:6]
                        identifier = (row['성명'], birthdate)
                        if identifier in processed_individuals:
                            continue
                        processed_individuals.add(identifier)
                        name = row['성명']
                        if name in name_count_map:
                            if identifier in name_count_map[name]:
                                name_count_map[name][identifier] += 1
                            else:
                                name_count_map[name][identifier] = 2
                            unique_name = f"{name}{name_count_map[name][identifier]}"
                        else:
                            name_count_map[name] = {identifier: 1}
                            unique_name = name
                        if tgt_row_idx_PS_HI_edi >= 45:
                            self.sheets['PS_HI_edi'].insert_rows(tgt_row_idx_PS_HI_edi)
                            self.copy_format_row(self.sheets['PS_HI_edi'], tgt_row_idx_PS_HI_edi - 1, tgt_row_idx_PS_HI_edi)
                        for col_name in HI_columns:
                            if col_name in row and col_name in tgt_columns:
                                col_idx = tgt_columns[col_name]
                                value = convert_to_number(row[col_name])
                                self.sheets['PS_HI_edi'].cell(row=tgt_row_idx_PS_HI_edi, column=col_idx, value=value)
                        self.sheets['PS_HI_edi'].cell(row=tgt_row_idx_PS_HI_edi, column=tgt_columns['성명'], value=unique_name)
                        tgt_row_idx_PS_HI_edi += 1
                        self.tgt_row_idx_PS_HI_edi = tgt_row_idx_PS_HI_edi

            def copy_health_insurance_data(df):
                tgt_row_idx_HI_edi = self.tgt_row_idx_HI_edi
                tgt_columns = {cell.value: cell.column for cell in self.sheets['HI_edi'][1]}
                processed_individuals = set()
                name_count_map = {}
                for idx, row in df.dropna(subset=['성명']).iterrows():
                    name = row['성명']
                    birthdate = None
                    if '주민등록번호' in row:
                        birthdate = row['주민등록번호'][:6]
                    elif '생년월일' in row:
                        birthdate = row['생년월일'].replace('-', '')
                    identifier = (name, birthdate)
                    if identifier in processed_individuals:
                        continue
                    processed_individuals.add(identifier)
                    if name in name_count_map:
                        if identifier in name_count_map[name]:
                            name_count_map[name][identifier] += 1
                        else:
                            name_count_map[name][identifier] = 2
                        unique_name = f"{name}{name_count_map[name][identifier]}"
                    else:
                        name_count_map[name] = {identifier: 1}
                        unique_name = name
                    if tgt_row_idx_HI_edi >= 45:
                        self.sheets['HI_edi'].insert_rows(tgt_row_idx_HI_edi)
                        self.copy_format_row(self.sheets['HI_edi'], tgt_row_idx_HI_edi - 1, tgt_row_idx_HI_edi)
                    for col_name in HI_columns:
                        col_idx = tgt_columns.get(col_name)
                        if col_idx and col_name in row:
                            value = convert_to_number(row[col_name])
                            self.sheets['HI_edi'].cell(row=tgt_row_idx_HI_edi, column=col_idx, value=value)
                    self.sheets['HI_edi'].cell(row=tgt_row_idx_HI_edi, column=tgt_columns['성명'], value=unique_name)
                    tgt_row_idx_HI_edi += 1
                    self.tgt_row_idx_HI_edi = tgt_row_idx_HI_edi

            def copy_np_insurance_data(df):
                tgt_row_idx_NP_edi = self.tgt_row_idx_NP_edi
                tgt_columns = {cell.value: cell.column for cell in self.sheets['NP_edi'][1]}
                processed_individuals = set()
                name_count_map = {}
                for idx, row in df.dropna(subset=['성명']).iterrows():
                    name = row['성명']
                    birthdate = None
                    if '주민등록번호' in row:
                        birthdate = row['주민등록번호'][:6]
                    elif '생년월일' in row:
                        birthdate = row['생년월일'].replace('-', '')
                    identifier = (name, birthdate)
                    if identifier in processed_individuals:
                        continue
                    processed_individuals.add(identifier)
                    if name in name_count_map:
                        if identifier in name_count_map[name]:
                            name_count_map[name][identifier] += 1
                        else:
                            name_count_map[name][identifier] = 2
                        unique_name = f"{name}{name_count_map[name][identifier]}"
                    else:
                        name_count_map[name] = {identifier: 1}
                        unique_name = name
                    if tgt_row_idx_NP_edi >= 28:
                        self.sheets['NP_edi'].insert_rows(tgt_row_idx_NP_edi)
                        self.copy_format_row(self.sheets['NP_edi'], tgt_row_idx_NP_edi - 1, tgt_row_idx_NP_edi)
                    for col_name in NP_columns:
                        col_idx = tgt_columns.get(col_name)
                        if col_idx and col_name in row:
                            value = convert_to_number(row[col_name])
                            self.sheets['NP_edi'].cell(row=tgt_row_idx_NP_edi, column=col_idx, value=value)
                    self.sheets['NP_edi'].cell(row=tgt_row_idx_NP_edi, column=tgt_columns['성명'], value=unique_name)
                    tgt_row_idx_NP_edi += 1
                    self.tgt_row_idx_NP_edi = tgt_row_idx_NP_edi

            def copy_ep_insurance_data(df):
                tgt_row_idx_EP_edi = self.tgt_row_idx_EP_edi
                tgt_columns = {cell.value: cell.column for cell in self.sheets['EP_edi'][1]}
                processed_individuals = set()
                name_count_map = {}
                for idx, row in df.dropna(subset=['근로자명']).iterrows():
                    name = row['근로자명']
                    birthdate = None
                    if '주민등록번호' in row:
                        birthdate = row['주민등록번호'][:6]
                    elif '생년월일' in row:
                        birthdate = row['생년월일'].replace('-', '')
                    identifier = (name, birthdate)
                    if identifier in processed_individuals:
                        continue
                    processed_individuals.add(identifier)
                    if name in name_count_map:
                        if identifier in name_count_map[name]:
                            name_count_map[name][identifier] += 1
                        else:
                            name_count_map[name][identifier] = 2
                        unique_name = f"{name}{name_count_map[name][identifier]}"
                    else:
                        name_count_map[name] = {identifier: 1}
                        unique_name = name
                    if tgt_row_idx_EP_edi >= 28:
                        self.sheets['EP_edi'].insert_rows(tgt_row_idx_EP_edi)
                        self.copy_format_row(self.sheets['EP_edi'], tgt_row_idx_EP_edi - 1, tgt_row_idx_EP_edi)
                    for col_name in EP_columns:
                        col_idx = tgt_columns.get(col_name)
                        value = None
                        if col_name == '근로자실업급여보험료' and len(row) > 22:
                            value = row.iloc[22]
                        elif col_name == '사업주실업급여보험료' and len(row) > 23:
                            value = row.iloc[23]
                        elif col_name == '사업주고안직능보험료' and len(row) > 24:
                            value = row.iloc[24]
                        elif col_name in row and col_name in tgt_columns:
                            value = convert_to_number(row[col_name])
                        if col_idx and value is not None:
                            self.sheets['EP_edi'].cell(row=tgt_row_idx_EP_edi, column=col_idx, value=value)
                    self.sheets['EP_edi'].cell(row=tgt_row_idx_EP_edi, column=tgt_columns['근로자명'], value=unique_name)
                    tgt_row_idx_EP_edi += 1
                    self.tgt_row_idx_EP_edi = tgt_row_idx_EP_edi

            def copy_ia_insurance_data(df):
                tgt_row_idx_IA_edi = self.tgt_row_idx_IA_edi
                tgt_columns = {cell.value: cell.column for cell in self.sheets['IA_edi'][1]}
                processed_individuals = set()
                name_count_map = {}
                for idx, row in df.dropna(subset=['근로자명']).iterrows():
                    name = row['근로자명']
                    birthdate = None
                    if '주민등록번호' in row:
                        birthdate = row['주민등록번호'][:6]
                    elif '생년월일' in row:
                        birthdate = row['생년월일'].replace('-', '')
                    identifier = (name, birthdate)
                    if identifier in processed_individuals:
                        continue
                    processed_individuals.add(identifier)
                    if name in name_count_map:
                        if identifier in name_count_map[name]:
                            name_count_map[name][identifier] += 1
                        else:
                            name_count_map[name][identifier] = 2
                        unique_name = f"{name}{name_count_map[name][identifier]}"
                    else:
                        name_count_map[name] = {identifier: 1}
                        unique_name = name
                    if tgt_row_idx_IA_edi >= 25:
                        self.sheets['IA_edi'].insert_rows(tgt_row_idx_IA_edi)
                        self.copy_format_row(self.sheets['IA_edi'], tgt_row_idx_IA_edi - 1, tgt_row_idx_IA_edi)
                    for col_name in IA_columns:
                        col_idx = tgt_columns.get(col_name)
                        if col_idx and col_name in row:
                            value = convert_to_number(row[col_name])
                            self.sheets['IA_edi'].cell(row=tgt_row_idx_IA_edi, column=col_idx, value=value)
                    self.sheets['IA_edi'].cell(row=tgt_row_idx_IA_edi, column=tgt_columns['근로자명'], value=unique_name)
                    tgt_row_idx_IA_edi += 1
                    self.tgt_row_idx_IA_edi = tgt_row_idx_IA_edi

            for src_file in src_files[:-1]:
                try:
                    if src_file:
                        df_src = pd.read_excel(src_file, index_col=None)
                        copy_data(df_src, is_original_file=True)
                except Exception as e:
                    continue

            try:
                if PS_health_insurance_file:
                    df_PS_health_insurance = pd.read_excel(PS_health_insurance_file)
                    copy_PS_health_insurance_data(df_PS_health_insurance)
            except Exception as e:
                pass

            try:
                if health_insurance_file:
                    df_health_insurance = pd.read_excel(health_insurance_file)
                    copy_health_insurance_data(df_health_insurance)
            except Exception as e:
                pass

            try:
                if np_insurance_file:
                    df_np_insurance = pd.read_excel(np_insurance_file)
                    copy_np_insurance_data(df_np_insurance)
            except Exception as e:
                pass

            try:
                if ep_insurance_file:
                    df_ep_insurance = pd.read_excel(ep_insurance_file, header=1)
                    copy_ep_insurance_data(df_ep_insurance)
            except Exception as e:
                pass

            try:
                if ia_insurance_file:
                    df_ia_insurance = pd.read_excel(ia_insurance_file)
                    copy_ia_insurance_data(df_ia_insurance)
            except Exception as e:
                pass

            try:
                if src_files[-1] and self.include_income_data.get():
                    df_income = pd.read_excel(src_files[-1])
                    copy_income_data(df_income)
            except Exception as e:
                pass

            self.apply_conditional_formatting(total_i_sheet)
            self.apply_conditional_formatting_zero(ps_health_sheet, ['L7:L102', 'C109'])
            self.apply_conditional_formatting_zero(total_i_sheet, ['R7:R87', 'V7:V87', 'AD7:AD87', 'AF7:AF87'])
            new_tgt_file = tgt_file.replace('.xlsm', f'_{current_date.year}년 {current_date.month}월.xlsm')
            temp_file = "temp_file.xlsm"
            for sheet_name in new_wb.sheetnames:
                sheet = new_wb[sheet_name]
                if sheet_name in ["(EDI) 공무원 건강보험", "(EDI) 비공무원 건강보험", "(EDI) 국민연금", "(EDI) 고용보험", "(EDI) 산재보험"]:
                    sheet.sheet_properties.tabColor = "FFFFEF99"
                elif sheet_name == "(급여대장) 월급여+기타4+기타1":
                    sheet.sheet_properties.tabColor = "FFC0CDEF"
                elif sheet_name == "사용법":
                    sheet.sheet_properties.tabColor = "FFFF0000"
                else:
                    sheet.sheet_properties.tabColor = "FF9BE5C8"
            new_wb.save(new_tgt_file)

            def save_final_file_path(file_path):
                try:
                    settings_file = 'settings.json'
                    if os.path.exists(settings_file):
                        with open(settings_file, 'r', encoding='utf-8') as f:
                            settings_data = json.load(f)
                    else:
                        settings_data = {}
                    settings_data['file_path'] = file_path
                    with open(settings_file, 'w', encoding='utf-8') as f:
                        json.dump(settings_data, f, ensure_ascii=False, indent=4)
                except Exception as e:
                    print(f"파일 경로 저장 중 오류 발생: {e}")

            def on_password_input(password):
                try:
                    if password:
                        self.run_excel_macro(new_tgt_file)
                        self.set_password_on_excel(new_tgt_file, password)
                        save_final_file_path(new_tgt_file)
                        open_file = messagebox.askyesno('성공', f'자료 병합이 완료되었습니다! {new_tgt_file}\n완성된 파일을 열까요?')
                        if open_file:
                            os.startfile(new_tgt_file)
                    else:
                        self.run_excel_macro(new_tgt_file)
                        save_final_file_path(new_tgt_file)
                        open_file = messagebox.askyesno('경고', '비밀번호를 입력하지 않아 파일이 보호되지 않았습니다. 그래도 완성된 파일을 열까요?')
                        if open_file:
                            os.startfile(new_tgt_file)
                except Exception as e:
                    error_message = f'오류 발생: {e}\n{traceback.format_exc()}'
                    ErrorDialog(self, error_message)

            PasswordDialog(self, on_password_input)
        except IndexError:
            print("오류 발생: 인덱스가 범위를 벗어남.")
        except Exception as e:
            error_message = f'오류 발생: {e}\n{traceback.format_exc()}'
            ErrorDialog(self, error_message)

if __name__ == '__main__':
    root = tk.Tk()
    app = merge(root)
    app.pack(fill=tk.BOTH, expand=True)
    app.mainloop()
